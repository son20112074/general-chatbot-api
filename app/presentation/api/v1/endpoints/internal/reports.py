import os
import re
import tempfile
from datetime import date, datetime
from typing import Dict, Optional

from docx import Document as DocxDocument
from docx.shared import Pt
from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Query, status
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession
import asyncio

from app.core.database import get_db, get_db_session
from app.crons.report_export_job import _process_template, _running_templates
from app.domain.models.file import File as FileModel
from app.domain.models.report import Report, ReportStatusEnum
from app.domain.models.report_template import FileModeEnum, ReportTemplate
from app.infrastructure.services.report_service import ReportService
from app.presentation.api.dependencies import get_current_user
from app.presentation.api.v1.schemas.report import (
    MarkdownReportCreate,
    ReportDocumentItem,
    ReportResponse,
    ReportTemplateCreate,
    ReportTemplateResponse,
    ReportTemplateUpdate,
)

router = APIRouter()


def _serialize_report(report) -> ReportResponse:
    return ReportResponse(
        id=report.id,
        name=report.name,
        template_id=report.template_id,
        template_name=report.template.name if report.template else None,
        frequency=report.template.frequency.value if report.template and report.template.frequency else None,
        content=report.content,
        status=report.status.value,
        period_start=report.period_start,
        period_end=report.period_end,
        file_url=report.file_url,
        created_by=report.created_by,
        creator_name=report.creator.full_name if report.creator else None,
        created_at=report.created_at,
        document_count=len(report.documents),
        documents=[ReportDocumentItem.model_validate(d) for d in report.documents],
    )


def _assert_template_owner(template, user_id: int) -> None:
    if template.created_by != user_id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Not authorized to access this template")


def _assert_report_owner(report, user_id: int) -> None:
    if report.created_by != user_id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Not authorized to access this report")


# ── Markdown → DOCX helpers ───────────────────────────────────────────────────

_INLINE_MD = re.compile(
    r"\*\*\*(.+?)\*\*\*"
    r"|___(.+?)___"
    r"|\*\*(.+?)\*\*"
    r"|__(.+?)__"
    r"|\*(.+?)\*"
    r"|_(.+?)_"
    r"|`(.+?)`",
    re.DOTALL,
)


def _strip_md_inline(text: str) -> str:
    return _INLINE_MD.sub(lambda m: next(g for g in m.groups() if g is not None), text)


def _apply_inline_md(para, text: str) -> None:
    last = 0
    for m in _INLINE_MD.finditer(text):
        if m.start() > last:
            para.add_run(text[last : m.start()])
        g = m.groups()
        if g[0] or g[1]:
            run = para.add_run(g[0] or g[1])
            run.bold = True
            run.italic = True
        elif g[2] or g[3]:
            run = para.add_run(g[2] or g[3])
            run.bold = True
        elif g[4] or g[5]:
            run = para.add_run(g[4] or g[5])
            run.italic = True
        elif g[6]:
            run = para.add_run(g[6])
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        last = m.end()
    if last < len(text):
        para.add_run(text[last:])


def _build_docx_from_markdown(title: str, markdown_text: str) -> bytes:
    doc = DocxDocument()
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"Created: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    doc.add_paragraph("")

    lines = markdown_text.splitlines()
    i = 0
    in_code = False
    code_buf: list = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if in_code:
                para = doc.add_paragraph("\n".join(code_buf))
                for run in para.runs:
                    run.font.name = "Courier New"
                    run.font.size = Pt(10)
                in_code = False
                code_buf = []
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            doc.add_heading(_strip_md_inline(m.group(2)), level=min(len(m.group(1)), 4))
            i += 1
            continue

        if re.match(r"^(-{3,}|\*{3,}|_{3,})\s*$", line):
            doc.add_paragraph("")
            i += 1
            continue

        m = re.match(r"^\s*[-*+]\s+(.*)", line)
        if m:
            _apply_inline_md(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1
            continue

        m = re.match(r"^\s*\d+\.\s+(.*)", line)
        if m:
            _apply_inline_md(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue

        m = re.match(r"^>\s?(.*)", line)
        if m:
            run = doc.add_paragraph().add_run(m.group(1))
            run.italic = True
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        _apply_inline_md(doc.add_paragraph(), line)
        i += 1

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name
    try:
        doc.save(tmp_path)
        with open(tmp_path, "rb") as f:
            data = f.read()
    finally:
        os.unlink(tmp_path)
    return data


def _save_markdown_docx(content: bytes) -> str:
    from pathlib import Path
    out_dir = Path("static/downloads/reports")
    out_dir.mkdir(parents=True, exist_ok=True)
    filename = f"report_md_{datetime.now().strftime('%Y%m%d_%H%M%S%f')}.docx"
    (out_dir / filename).write_bytes(content)
    return f"static/downloads/reports/{filename}"


# ── Report Templates ──────────────────────────────────────────────────────────

@router.get("/templates", response_model=Dict)
async def list_templates(
    q: Optional[str] = Query(None, description="Search by template name"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    service = ReportService(db)
    result = await service.list_templates(created_by=current_user.user_id, q=q, page=page, page_size=page_size)
    return {
        "data": [ReportTemplateResponse.model_validate(t) for t in result["data"]],
        "total": result["total"],
        "page": result["page"],
        "page_size": result["page_size"],
    }


@router.post("/templates", response_model=ReportTemplateResponse, status_code=status.HTTP_201_CREATED)
async def create_template(
    data: ReportTemplateCreate,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    service = ReportService(db)
    try:
        template = await service.create_template(data, current_user.user_id)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    return ReportTemplateResponse.model_validate(template)


@router.get("/templates/{template_id}", response_model=ReportTemplateResponse)
async def get_template(
    template_id: int,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    service = ReportService(db)
    template = await service.get_template(template_id)
    if not template:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Template not found")
    _assert_template_owner(template, current_user.user_id)
    return ReportTemplateResponse.model_validate(template)


@router.put("/templates/{template_id}", response_model=ReportTemplateResponse)
async def update_template(
    template_id: int,
    data: ReportTemplateUpdate,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    service = ReportService(db)
    template = await service.get_template(template_id)
    if not template:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Template not found")
    _assert_template_owner(template, current_user.user_id)
    try:
        template = await service.update_template(template_id, data)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    return ReportTemplateResponse.model_validate(template)


@router.delete("/templates/{template_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_template(
    template_id: int,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    service = ReportService(db)
    template = await service.get_template(template_id)
    if not template:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Template not found")
    _assert_template_owner(template, current_user.user_id)
    await service.delete_template(template_id)


# ── Reports ───────────────────────────────────────────────────────────────────

async def _run_report_once_bg(template_id: int) -> None:
    if template_id in _running_templates:
        return

    today = date.today()
    async with get_db_session() as session:
        tpl_result = await session.execute(
            select(ReportTemplate).where(ReportTemplate.id == template_id)
        )
        template = tpl_result.scalar_one_or_none()
        if not template or not template.file_ids:
            return

        existing = await session.scalar(
            select(func.count()).select_from(Report)
            .where(Report.template_id == template_id)
            .where(Report.period_start == today)
            .where(Report.period_end == today)
            .where(Report.status.in_([ReportStatusEnum.COMPLETED, ReportStatusEnum.FAILED]))
        )
        if existing:
            return

        if template_id in _running_templates:
            return
        _running_templates.add(template_id)

        files_result = await session.execute(
            select(FileModel).where(
                FileModel.id.in_(template.file_ids),
                or_(FileModel.is_deleted == False, FileModel.is_deleted == None),
                FileModel.content.isnot(None),
                FileModel.content != "",
            )
        )
        files = list(files_result.scalars().all())
        if not files:
            _running_templates.discard(template_id)
            return

        try:
            await _process_template(session, template, files, today, today)
        finally:
            _running_templates.discard(template_id)


@router.post("/templates/{template_id}/run", status_code=status.HTTP_202_ACCEPTED, tags=["Reports"])
async def run_report_once(
    template_id: int,
    # background_tasks: BackgroundTasks,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    service = ReportService(db)
    template = await service.get_template(template_id)
    if not template:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Template not found")
    _assert_template_owner(template, current_user.user_id)
    if template.file_mode != FileModeEnum.SELECT:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Run once is only supported for select-mode templates")

    # background_tasks.add_task(_run_report_once_bg, template_id)
    asyncio.create_task(_run_report_once_bg(template_id))
    return {"message": "Report generation started"}


@router.post("/from-markdown", response_model=Dict, status_code=status.HTTP_200_OK, tags=["Reports"])
async def create_docx_from_markdown(
    data: MarkdownReportCreate,
    _: None = Depends(get_current_user),
):
    docx_bytes = _build_docx_from_markdown("Report", data.markdown)
    file_url = _save_markdown_docx(docx_bytes)
    return {"file_url": file_url}


@router.get("", response_model=Dict, tags=["Reports"])
async def list_reports(
    q: Optional[str] = Query(None, description="Search by report name"),
    status_filter: Optional[str] = Query(None, alias="status", description="compiling | completed | failed"),
    template_id: Optional[int] = Query(None),
    start_date: Optional[date] = Query(None, description="Filter by start date (report created_at)"),
    end_date: Optional[date] = Query(None, description="Filter by end date (report created_at)"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    service = ReportService(db)
    try:
        result = await service.list_reports(
            q=q,
            status=status_filter,
            template_id=template_id,
            created_by=current_user.user_id,
            start_date=start_date,
            end_date=end_date,
            page=page,
            page_size=page_size,
        )
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    return {
        "data": [_serialize_report(r) for r in result["data"]],
        "total": result["total"],
        "page": result["page"],
        "page_size": result["page_size"],
    }


@router.get("/{report_id}", response_model=ReportResponse, tags=["Reports"])
async def get_report(
    report_id: int,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    service = ReportService(db)
    report = await service.get_report(report_id)
    if not report:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Report not found")
    _assert_report_owner(report, current_user.user_id)
    return _serialize_report(report)



@router.delete("/{report_id}/documents", status_code=status.HTTP_204_NO_CONTENT, tags=["Reports"])
async def remove_documents(
    report_id: int,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    service = ReportService(db)
    report = await service.get_report(report_id)
    if not report:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Report not found")
    _assert_report_owner(report, current_user.user_id)
    await service.remove_documents(report_id)


@router.delete("/{report_id}", status_code=status.HTTP_204_NO_CONTENT, tags=["Reports"])
async def delete_report(
    report_id: int,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    service = ReportService(db)
    report = await service.get_report(report_id)
    if not report:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Report not found")
    _assert_report_owner(report, current_user.user_id)
    await service.delete_report(report_id)
